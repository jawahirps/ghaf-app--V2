"""
Ghaf App — Import CLAIMACTIVITY Excel exports into claim_activities table.
Also backfills real clinician_id (DHA-P-XXXXXX) onto the claims rows.

Supports TWO formats (auto-detected from header row):

FORMAT A — Al Noor (header row 1, data from row 2):
  0  Facility ID        1  Facility Name      2  Invoice No (claim_id)
  3  Receiver Id        4  Receiver Name      5  Payer ID
  6  Payer Name         7  Patient ID         8  Member Id
  9  Emirates Id       10  Encounter Start   11  Encounter End
 12  Transaction Date  13  Clinician ID      14  Clinician Name
 15  Prior Auth ID     16  CPT Type          17  Claim Activity Number
 18  CPT Code          19  Quantity          20  Initial Denial Code
 ...  Net Amount at 24

FORMAT B — Analytika Reports (header row 1, data from row 2):
  0  Facility Name      1  Transaction Ref    2  Receiver
  3  Receiver Name      4  Payer ID           5  Payer Name
  6  Patient ID         7  Member Id          8  Emirates Id
  9  Encounter Start   10  Encounter End     11  Transaction Date
 12  Clinician ID      13  Clinician Name    14  Prior Auth ID
 15  CPT Type          16  Claim Activity Number  17  CPT Code
 18  Quantity          ...  Net Amount (further right)

The key difference: Al Noor has "Facility ID" in col 0 then "Facility Name" in col 1.
Analytika starts directly with "Facility Name" in col 0.

Both formats use COLUMN HEADERS in row 1 and data from row 2.
This script detects the layout automatically via the header row.

USAGE (from backend/ directory with venv active):
    python -m app.import_claim_activity "/path/CLAIMACTIVITY_Q1 2024.xlsx" ...

Flags:
  --limit N    Only process first N rows per file (testing)
  --clear      Delete all existing EXCEL_HISTORICAL claim_activities before import
"""
from __future__ import annotations

import argparse
from datetime import datetime, date
from decimal import Decimal, InvalidOperation

import openpyxl
from psycopg2.extras import execute_values

from .database import SessionLocal, engine
from .import_claim_summary import parse_date, parse_decimal   # reuse helpers

BATCH_SIZE = 500

# ---------------------------------------------------------------------------
# Column name aliases — maps canonical field names to every known header variant
# ---------------------------------------------------------------------------
COL_ALIASES: dict[str, list[str]] = {
    "facility_name":      ["Facility Name", "Facility"],
    "facility_license_id":["Facility ID", "DHA Facility ID", "License No", "Facility Code"],
    "claim_id":           ["Transaction Ref", "Invoice No", "Invoice Number", "Claim Number", "TransactionRef"],
    "receiver_id":        ["Receiver Id", "Receiver"],
    "receiver_name":      ["Receiver Name"],
    "payer_id":           ["Payer ID", "Payer"],
    "payer_name":         ["Payer Name"],
    "patient_id":         ["Patient Id", "Patient ID"],
    "member_id":          ["Member Id", "Member ID"],
    "emirates_id":        ["Emirates Id", "Emirates ID"],
    "encounter_start":    ["Encounter Start", "Treatment Date", "Date Of Service"],
    "encounter_end":      ["Encounter End"],
    "transaction_date":   ["Transaction Date", "Activity Date"],
    "clinician_id":       ["Clinician ID", "Clinician Id", "Provider ID"],
    "clinician_name":     ["Clinician Name", "Provider Name"],
    "auth_id":            ["Prior Auth ID", "Prior Authorization ID", "Authorization Id"],
    "activity_type_raw":  ["CPT Type", "Activity Type", "Service Type"],
    "sequence_no_raw":    ["Claim Activity Number", "Activity Number", "Sequence No"],
    "activity_code":      ["CPT Code", "Activity Code", "Service Code",
                           "Procedure Code", "Treatment Code", "HCPCS Code",
                           "CPT / HCPCS Code", "CPT/HCPCS Code", "Activity Identification",
                           "Claim Activity Code", "Act Code", "Code"],
    "quantity":           ["Quantity", "Qty"],
    "initial_denial_code":["Initial Denial Code"],
    "initial_denial_desc":["Initial Denial Comment"],
    "net_amount":         ["Net Amount", "Net Amt", "Net Amt - Initial Sub"],
    "remitted_amt":       ["Remitted Amt", "Remitted Amount"],
    "initial_rejected":   ["Initial Rejected", "Initial Sub Rejected Amt"],
    "rejected_amt":       ["Rejected Amt", "Rejected Amount"],
    "payment_ref":        ["Payment Ref", "Payment Reference"],
    "settlement_date":    ["Settlement Date"],
    "denial_code":        ["Denial Code"],
    "denial_comment":     ["Denial Comment", "Denial Description"],
    "diagnosis":          ["Diagnosis"],
    "claim_year":         ["Claim Year"],
    "claim_month":        ["Claim Month"],
}

_HEADER_LOOKUP: dict[str, str] = {}
for _field, _aliases in COL_ALIASES.items():
    for _alias in _aliases:
        _HEADER_LOOKUP[_alias.lower().strip()] = _field

# ---------------------------------------------------------------------------
# activity_type mapping
# ---------------------------------------------------------------------------
ACTIVITY_TYPE_MAP = {
    "service code": "Service",
    "service":      "Service",
    "cpt":          "CPT",
    "hcpcs":        "HCPCS",
    "drug":         "Drug",
    "dental":       "Dental",
}

ACTIVITY_SQL = """
INSERT INTO claim_activities (
    claim_id_fk, facility_id,
    activity_type, activity_code, activity_description,
    quantity, gross, net, patient_share,
    activity_start, authorization_id, clinician_id, sequence_no
) VALUES %s
"""

CLINICIAN_UPDATE_SQL = """
UPDATE claims SET clinician_id = %s, updated_at = NOW()
WHERE id = %s AND (clinician_id IS NULL OR clinician_id LIKE 'CLN%%')
"""

# Minimal claim record created when CLAIMSUMMARY hasn't been imported yet
CLAIM_UPSERT_SQL = """
INSERT INTO claims (
    facility_id, source_channel, claim_id, encounter_id,
    encounter_type, encounter_start,
    payer_id, payer_name, receiver_id, receiver_name,
    member_id, clinician_id, clinician_name,
    gross_amount, net_amount, patient_share_total,
    submission_date, claim_status, resubmission_count
) VALUES (
    %(facility_id)s, 'EXCEL_HISTORICAL', %(claim_id)s, %(claim_id)s,
    %(encounter_type)s, %(encounter_start)s,
    %(payer_id)s, %(payer_name)s, %(receiver_id)s, %(receiver_name)s,
    %(member_id)s, %(clinician_id)s, %(clinician_name)s,
    0, 0, 0,
    %(encounter_start)s, 'submitted', 0
)
ON CONFLICT (facility_id, claim_id) DO NOTHING
RETURNING id
"""


# ---------------------------------------------------------------------------
# Format detection
# ---------------------------------------------------------------------------
def detect_col_map(ws) -> dict[str, int]:
    """
    Reads row 1 as the header row and builds a {field_name: col_index} map.
    Both Al Noor and Analytika formats have headers in row 1.
    Falls back to partial-match for activity_code if exact match fails.
    """
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return {}
    col_map: dict[str, int] = {}
    raw_headers: list[tuple[int, str]] = []
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        normalized = str(cell).lower().strip()
        raw_headers.append((idx, normalized))
        field = _HEADER_LOOKUP.get(normalized)
        if field and field not in col_map:
            col_map[field] = idx

    # Fuzzy fallback: if activity_code not matched, find any header containing
    # "cpt", "activity code", "procedure code", or "act id" (avoid matching
    # "activity type", "activity number" which have their own fields)
    if "activity_code" not in col_map:
        for idx, h in raw_headers:
            if (("cpt" in h or "procedure code" in h or "act id" in h)
                    and "activity_code" not in col_map
                    and "type" not in h and "number" not in h and "name" not in h):
                col_map["activity_code"] = idx
                print(f"  [fuzzy] Mapped activity_code from header: '{h}' at col {idx}", flush=True)
                break
        # Last resort: "activity code" substring
        if "activity_code" not in col_map:
            for idx, h in raw_headers:
                if "activity code" in h and "activity_code" not in col_map:
                    col_map["activity_code"] = idx
                    print(f"  [fuzzy] Mapped activity_code from header: '{h}' at col {idx}", flush=True)
                    break

    return col_map


# ---------------------------------------------------------------------------
# Lookups
# ---------------------------------------------------------------------------
def load_facility_lookups(raw_conn):
    """
    Returns three dicts:
      exact   — facility_name (exact) → id
      normed  — facility_name (lower+stripped) → id
      by_lic  — license_id (DHA-F-XXXXXXX, lower) → id
    """
    cur = raw_conn.cursor()
    cur.execute("SELECT id, facility_name, license_id FROM facilities")
    rows = cur.fetchall()
    cur.close()
    exact   = {name: fid for fid, name, _lic in rows}
    normed  = {name.lower().strip(): fid for fid, name, _lic in rows}
    by_lic  = {lic.lower().strip(): fid for fid, _name, lic in rows if lic}
    print(f"Facility lookup loaded: {len(exact)} facilities — {list(exact.keys())}")
    print(f"  License IDs in DB: {list(by_lic.keys())}")
    return exact, normed, by_lic


def load_claim_lookup(raw_conn) -> dict[tuple, tuple]:
    cur = raw_conn.cursor()
    cur.execute(
        "SELECT id, facility_id, claim_id, clinician_id "
        "FROM claims WHERE source_channel = 'EXCEL_HISTORICAL'"
    )
    lookup: dict[tuple, tuple] = {}
    for row in cur.fetchall():
        db_id, fac_id, claim_id, clin_id = row
        lookup[(fac_id, claim_id)] = (db_id, clin_id)
    cur.close()
    print(f"Claim lookup loaded: {len(lookup):,} entries.")
    return lookup


# ---------------------------------------------------------------------------
# Row → activity tuple (uses col_map instead of hardcoded indices)
# ---------------------------------------------------------------------------
def get_col(row: tuple, col_map: dict, field: str, default=None):
    idx = col_map.get(field)
    if idx is None or idx >= len(row):
        return default
    return row[idx]


def map_activity_type(raw: str) -> str:
    if not raw:
        return "Service"
    return ACTIVITY_TYPE_MAP.get(raw.strip().lower(), "Service")


def build_activity_tuple(row: tuple, col_map: dict, claim_db_id: int, facility_id: int):
    try:
        activity_code = str(get_col(row, col_map, "activity_code") or "").strip()
        if not activity_code:
            return None

        activity_start = (
            parse_date(get_col(row, col_map, "encounter_start"))
            or parse_date(get_col(row, col_map, "transaction_date"))
        )
        if activity_start is None:
            return None

        activity_type = map_activity_type(str(get_col(row, col_map, "activity_type_raw") or ""))
        quantity = parse_decimal(get_col(row, col_map, "quantity")) or Decimal("1")
        net = parse_decimal(get_col(row, col_map, "net_amount"))

        seq_raw = get_col(row, col_map, "sequence_no_raw")
        try:
            seq = int(seq_raw or 1)
            seq = max(1, min(seq, 32767))  # clamp to smallint range
        except (ValueError, TypeError):
            seq = 1

        clinician_id_val = str(get_col(row, col_map, "clinician_id") or "").strip() or None
        auth_id_val = str(get_col(row, col_map, "auth_id") or "").strip() or None

        return (
            claim_db_id,
            facility_id,
            activity_type,
            activity_code[:50],
            None,                                                  # activity_description
            quantity,
            net,                                                   # gross = net
            net,
            Decimal("0"),                                          # patient_share
            activity_start,
            auth_id_val[:100] if auth_id_val else None,
            clinician_id_val[:50] if clinician_id_val else None,
            seq,
        )
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Flush
# ---------------------------------------------------------------------------
def flush(raw_conn, activity_tuples: list, clinician_updates: list) -> None:
    if not activity_tuples and not clinician_updates:
        return
    cur = raw_conn.cursor()
    try:
        if activity_tuples:
            execute_values(cur, ACTIVITY_SQL, activity_tuples)
        if clinician_updates:
            cur.executemany(CLINICIAN_UPDATE_SQL, clinician_updates)
        raw_conn.commit()
    except Exception:
        raw_conn.rollback()
        raise
    finally:
        cur.close()
    activity_tuples.clear()
    clinician_updates.clear()


# ---------------------------------------------------------------------------
# Process one file
# ---------------------------------------------------------------------------
def _clinician_id_from_name(name: str | None) -> str:
    if not name or not str(name).strip():
        return "UNKNOWN"
    digest = abs(hash(str(name).strip().upper())) % (10 ** 8)
    return f"CLN{digest:08d}"


def _ensure_claim(raw_conn, fac_id: int, claim_id_str: str,
                  row: tuple, col_map: dict,
                  claim_lookup: dict) -> tuple[int, str] | None:
    """
    If the claim doesn't exist in claim_lookup, create a minimal record and
    add it to the lookup so subsequent rows for the same claim use the same id.
    Returns (db_id, clinician_id) or None on error.
    """
    cur = raw_conn.cursor()
    encounter_start = (
        parse_date(get_col(row, col_map, "encounter_start"))
        or parse_date(get_col(row, col_map, "transaction_date"))
    )
    if encounter_start is None:
        cur.close()
        return None

    from datetime import datetime as _dt
    clin_name = str(get_col(row, col_map, "clinician_name") or "").strip() or None
    clin_id_raw = str(get_col(row, col_map, "clinician_id") or "").strip() or None
    clin_id = clin_id_raw if (clin_id_raw and clin_id_raw.upper().startswith("DHA-P-")) \
              else _clinician_id_from_name(clin_name)

    params = {
        "facility_id":    fac_id,
        "claim_id":       claim_id_str,
        "encounter_type": str(get_col(row, col_map, "activity_type_raw") or "").strip() or None,
        "encounter_start": _dt.combine(encounter_start, _dt.min.time()),
        "payer_id":       str(get_col(row, col_map, "payer_id") or "UNKNOWN").strip()[:50],
        "payer_name":     str(get_col(row, col_map, "payer_name") or "").strip() or None,
        "receiver_id":    str(get_col(row, col_map, "receiver_id") or "").strip() or None,
        "receiver_name":  str(get_col(row, col_map, "receiver_name") or "").strip() or None,
        "member_id":      str(get_col(row, col_map, "member_id") or claim_id_str).strip()[:100],
        "clinician_id":   clin_id[:50],
        "clinician_name": clin_name,
    }
    try:
        cur.execute(CLAIM_UPSERT_SQL, params)
        row_ret = cur.fetchone()
        if row_ret:
            db_id = row_ret[0]
        else:
            # ON CONFLICT DO NOTHING — fetch existing
            cur.execute(
                "SELECT id, clinician_id FROM claims WHERE facility_id=%s AND claim_id=%s",
                (fac_id, claim_id_str),
            )
            fetched = cur.fetchone()
            if not fetched:
                cur.close()
                return None
            db_id, clin_id = fetched[0], fetched[1]
        raw_conn.commit()
        claim_lookup[(fac_id, claim_id_str)] = (db_id, clin_id)
        return db_id, clin_id
    except Exception as e:
        raw_conn.rollback()
        print(f"  ⚠ Could not create claim {claim_id_str}: {e}", flush=True)
        return None
    finally:
        cur.close()


def process_file(raw_conn, path: str, claim_lookup: dict,
                 facility_name_lookup: dict, facility_name_lookup_normed: dict,
                 facility_license_lookup: dict,
                 limit=None):
    print(f"Opening {path} ...", flush=True)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    col_map = detect_col_map(ws)
    if not col_map or "claim_id" not in col_map:
        print(f"  WARNING: Could not detect required column (claim_id) in {path}, skipping.", flush=True)
        wb.close()
        return 0, 0

    if "facility_name" not in col_map and "facility_license_id" not in col_map:
        print(f"  WARNING: No facility column detected in {path}, skipping.", flush=True)
        wb.close()
        return 0, 0

    print(f"  Columns detected: facility_name@{col_map.get('facility_name','–')}, "
          f"facility_license_id@{col_map.get('facility_license_id','–')}, "
          f"claim_id@{col_map['claim_id']}, "
          f"activity_code@{col_map.get('activity_code', '?')}, "
          f"net_amount@{col_map.get('net_amount', '?')}", flush=True)

    activity_tuples: list = []
    clinician_updates: list = []
    total = 0
    skipped = 0
    created_claims = 0
    unknown_facilities: set = set()   # for diagnostic output

    for row in ws.iter_rows(min_row=2, values_only=True):   # row 1 = headers
        if limit is not None and total >= limit:
            break
        if not row or all(v is None for v in row[:4]):
            continue

        facility_name       = str(get_col(row, col_map, "facility_name") or "").strip()
        facility_license_id = str(get_col(row, col_map, "facility_license_id") or "").strip()
        claim_id_str        = str(get_col(row, col_map, "claim_id") or "").strip()

        if not claim_id_str:
            skipped += 1
            continue

        # 1) Match by DHA license ID (most reliable — col A in Al Noor format)
        fac_id = facility_license_lookup.get(facility_license_id.lower()) if facility_license_id else None
        # 2) Exact facility name match
        if fac_id is None and facility_name:
            fac_id = facility_name_lookup.get(facility_name)
        # 3) Case-insensitive / stripped name match
        if fac_id is None and facility_name:
            fac_id = facility_name_lookup_normed.get(facility_name.lower().strip())
        if fac_id is None:
            label = facility_license_id or facility_name or "<unknown>"
            unknown_facilities.add(label)
            skipped += 1
            continue

        # Find claim — create minimal record if missing
        match = claim_lookup.get((fac_id, claim_id_str))
        if match is None:
            result = _ensure_claim(raw_conn, fac_id, claim_id_str, row, col_map, claim_lookup)
            if result is None:
                skipped += 1
                continue
            match = claim_lookup[(fac_id, claim_id_str)]
            created_claims += 1

        db_id, current_clinician = match

        act = build_activity_tuple(row, col_map, db_id, fac_id)
        if act is None:
            skipped += 1
            continue

        activity_tuples.append(act)

        # Backfill real DHA-P-XXXXXX clinician ID if available
        new_clin = str(get_col(row, col_map, "clinician_id") or "").strip()
        if new_clin and new_clin.upper().startswith("DHA-P-"):
            if not current_clinician or current_clinician.startswith("CLN"):
                clinician_updates.append((new_clin, db_id))
                claim_lookup[(fac_id, claim_id_str)] = (db_id, new_clin)

        total += 1
        if len(activity_tuples) >= BATCH_SIZE:
            flush(raw_conn, activity_tuples, clinician_updates)
            print(f"  ... {total} rows processed", flush=True)

    flush(raw_conn, activity_tuples, clinician_updates)
    wb.close()
    if unknown_facilities:
        print(f"  ⚠ Unknown facility names (not in DB): {sorted(unknown_facilities)[:10]}", flush=True)
    if created_claims:
        print(f"  ℹ Created {created_claims} minimal claim records (no prior CLAIMSUMMARY).", flush=True)
    print(f"Done {path}: {total} activities imported, {skipped} skipped.", flush=True)
    return total, skipped


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Import CLAIMACTIVITY_*.xlsx into claim_activities.")
    parser.add_argument("files", nargs="+", help="CLAIMACTIVITY_*.xlsx file paths")
    parser.add_argument("--limit", type=int, default=None, help="Max rows per file (testing)")
    parser.add_argument("--clear", action="store_true",
                        help="Delete existing EXCEL_HISTORICAL activities before import")
    args = parser.parse_args()

    raw_conn = engine.raw_connection()
    try:
        if args.clear:
            cur = raw_conn.cursor()
            cur.execute(
                "DELETE FROM claim_activities WHERE claim_id_fk IN "
                "(SELECT id FROM claims WHERE source_channel = 'EXCEL_HISTORICAL')"
            )
            deleted = cur.rowcount
            raw_conn.commit()
            cur.close()
            print(f"Cleared {deleted:,} existing EXCEL_HISTORICAL activities.")

        facility_name_lookup, facility_name_lookup_normed, facility_license_lookup = load_facility_lookups(raw_conn)
        claim_lookup = load_claim_lookup(raw_conn)

        grand_total = 0
        for path in args.files:
            total, _ = process_file(
                raw_conn, path, claim_lookup,
                facility_name_lookup, facility_name_lookup_normed,
                facility_license_lookup,
                args.limit,
            )
            grand_total += total

        print(f"\nAll done. {grand_total:,} activities imported across {len(args.files)} file(s).")
    finally:
        raw_conn.close()


if __name__ == "__main__":
    main()
